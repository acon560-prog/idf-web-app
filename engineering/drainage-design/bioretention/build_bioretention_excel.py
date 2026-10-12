#!/usr/bin/env python3
"""
Bio-rétention — pack de dimensionnement v1 (formules Excel live).

Produit vendable / livrable:
  - surface filtrante requise Af
  - volume de captage Vcap
  - temps de vidange (ponding)
  - bilan de couches (mulch / média / réservoir)
  - note Méthode + checklist CSA W200 (rappels, pas le texte de la norme)

Références de pratique (Canada): STEP LID wiki sizing, esprit CSA W200:18.
Jaune = entrées · vert = résultats · bleu = intermédiaires.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
OUT = HERE / "Bioretention_Dimensionnement_v1.xlsx"

YELLOW = PatternFill("solid", fgColor="FFF59D")
GREEN = PatternFill("solid", fgColor="C8E6C9")
BLUE = PatternFill("solid", fgColor="BBDEFB")
ORANGE = PatternFill("solid", fgColor="FFE0B2")
HDR = PatternFill("solid", fgColor="1B5E4B")
HDR_FONT = Font(bold=True, color="FFFFFF")
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def hdr_row(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HDR
        cell.font = HDR_FONT
        cell.border = THIN


def yellow(ws, r, c, val, fmt=None):
    cell = ws.cell(r, c, val)
    cell.fill = YELLOW
    cell.border = THIN
    if fmt:
        cell.number_format = fmt
    return cell


def green(ws, r, c, val, fmt=None):
    cell = ws.cell(r, c, val)
    cell.fill = GREEN
    cell.border = THIN
    if fmt:
        cell.number_format = fmt
    return cell


def blue(ws, r, c, val, fmt=None):
    cell = ws.cell(r, c, val)
    cell.fill = BLUE
    cell.border = THIN
    if fmt:
        cell.number_format = fmt
    return cell


def build() -> Path:
    wb = Workbook()

    # ---------- Parametres ----------
    ws = wb.active
    ws.title = "Parametres"
    ws["A1"] = "BIO-RÉTENTION — Dimensionnement v1 (formules Excel)"
    ws["A1"].font = Font(bold=True, size=14, color="1B5E4B")
    ws.merge_cells("A1:F1")
    ws["A2"] = (
        "But: estimer la surface filtrante Af et vérifier la vidange pour un volume de captage "
        "(qualité d'eau / petite pluie). Jaune = entrées. Vert = résultats. "
        "Adapter aux exigences municipales + CSA W200."
    )
    ws.merge_cells("A2:F3")
    ws["A2"].alignment = Alignment(wrap_text=True)

    ws["A5"] = "1) Bassin versant (jaune)"
    ws["A5"].font = Font(bold=True)
    rows_bv = [
        (6, "Aire drainée Ad (m2)", 2500, "0"),
        (7, "Coefficient de ruissellement C (-)", 0.70, "0.00"),
        (8, "Profondeur de pluie de captage Rd (mm)", 25, "0.0"),
        (9, "Facteur de captage η (-)", 0.90, "0.00"),
    ]
    for r, lab, val, fmt in rows_bv:
        ws.cell(r, 1, lab)
        yellow(ws, r, 2, val, fmt)
    ws["C8"] = "Souvent 25 mm (première chasse / WQ) — confirmer ville"
    ws["C9"] = "Fraction du volume théorique réellement ciblée (0.8–1.0)"

    ws["A11"] = "2) Géométrie & sols (jaune)"
    ws["A11"].font = Font(bold=True)
    rows_geo = [
        (12, "Profondeur d'ennoiement max dp (m)", 0.30, "0.00"),
        (13, "Profondeur média filtrant df (m)", 0.60, "0.00"),
        (14, "Porosité effective média nf (-)", 0.35, "0.00"),
        (15, "Profondeur réservoir/gravier dr (m)", 0.30, "0.00"),
        (16, "Porosité effective réservoir nr (-)", 0.40, "0.00"),
        (17, "Profondeur mulch dm (m)", 0.075, "0.000"),
        (18, "Porosité mulch nm (-)", 0.20, "0.00"),
        (19, "Taux d'infiltration de conception f' (mm/h)", 25, "0.0"),
        (20, "Temps de vidange max admis t_max (h)", 48, "0"),
        (21, "Sous-drain présent? (1=oui, 0=non)", 1, "0"),
    ]
    for r, lab, val, fmt in rows_geo:
        ws.cell(r, 1, lab)
        yellow(ws, r, 2, val, fmt)
    ws["C13"] = "0.30 pelouse/toit · 0.60 arbustes · 1.00 arbres (STEP typique)"
    ws["C19"] = "f' de conception (facteur de sécurité sur essai in situ)"
    ws["C21"] = "1 = avec underdrain; 0 = infiltration seule"

    ws["A23"] = "3) Résultats (formules live)"
    ws["A23"].font = Font(bold=True, size=12, color="1B5E4B")

    # Vcap = η * C * Rd(mm)/1000 * Ad
    ws["A24"] = "Volume de captage Vcap (m3)"
    green(ws, 24, 2, "=B9*B7*(B8/1000)*B6", "0.00")
    ws["C24"] = "Vcap = η × C × (Rd/1000) × Ad"

    # Effective storage depth ds (m)
    # ds = dp + dm*nm + df*nf + dr*nr   (simplified storage available per m2 of filter)
    ws["A25"] = "Profondeur de stockage effective ds (m)"
    blue(ws, 25, 2, "=B12+B17*B18+B13*B14+B15*B16", "0.000")
    ws["C25"] = "ds = dp + dm·nm + df·nf + dr·nr"

    # Af = Vcap / ds
    ws["A26"] = "Surface filtrante requise Af (m2)"
    green(ws, 26, 2, "=IF(B25<=0,\"ds=0\",B24/B25)", "0.0")
    ws["C26"] = "Af = Vcap / ds"

    # Ratio Af/Ad
    ws["A27"] = "Ratio Af / Ad (%)"
    green(ws, 27, 2, "=IF(B6<=0,0,100*B26/B6)", "0.0")
    ws["C27"] = "Ordre de grandeur typique souvent 5–20 % selon Rd et C"

    # Ponding drawdown time (h): t_p = dp / (f'/1000)   f' in mm/h → m/h
    ws["A28"] = "Temps de vidange de l'ennoiement tp (h)"
    green(ws, 28, 2, "=IF(B19<=0,\"f'=0\",B12/(B19/1000))", "0.0")
    ws["C28"] = "tp = dp / (f'/1000) — ennoiement de surface"

    # Media+reservoir drain estimate (simplified) through filter to underdrain/soil
    ws["A29"] = "Temps approximatif média→réservoir tm (h)"
    green(ws, 29, 2, "=IF(B19<=0,\"f'=0\",(B13*B14+B15*B16)/(B19/1000))", "0.0")
    ws["C29"] = "Approximatif: (df·nf + dr·nr) / (f'/1000)"

    ws["A30"] = "Vidange OK vs t_max?"
    green(ws, 30, 2, '=IF(OR(B28="f\'=0",B28=""),"",IF(B28<=B20,"OUI — tp ≤ t_max","NON — augmenter Af, dp↓ ou f\'"))')
    ws["B30"].fill = ORANGE

    ws["A31"] = "Longueur indicative L si largeur = 3 m"
    yellow(ws, 31, 2, 3.0, "0.0")
    ws["A32"] = "Longueur L (m) pour largeur B31"
    green(ws, 32, 2, "=IF(B31<=0,\"\",B26/B31)", "0.0")

    ws["A34"] = "Message produit"
    ws["A34"].font = Font(bold=True)
    ws["A35"] = (
        "Livrable: Af + Vcap + tp. Ce n'est PAS un design final CSA W200 — "
        "valider géotechnique (f'), plantation, prétraitement, trop-plein et exigences municipales."
    )
    ws.merge_cells("A35:F36")
    ws["A35"].alignment = Alignment(wrap_text=True)

    for col, w in zip("ABCDEF", [46, 14, 55, 12, 12, 12]):
        ws.column_dimensions[col].width = w

    # ---------- Couches ----------
    wc = wb.create_sheet("Couches")
    wc["A1"] = "Profil de couches — volumes pour Af calculée"
    wc["A1"].font = Font(bold=True, size=12, color="1B5E4B")
    wc["A2"] = "Af liée à Parametres!B26. Volumes = Af × profondeur × porosité (stockage)."
    for j, h in enumerate(["Couche", "Profondeur (m)", "Porosité n", "Stockage (m3)", "Volume solide+vide (m3)"], 1):
        wc.cell(4, j, h)
    hdr_row(wc, 4, 5)

    layers = [
        (5, "Ennoiement (ponding)", "=Parametres!B12", 1.0),
        (6, "Mulch", "=Parametres!B17", "=Parametres!B18"),
        (7, "Média filtrant", "=Parametres!B13", "=Parametres!B14"),
        (8, "Réservoir / gravier", "=Parametres!B15", "=Parametres!B16"),
    ]
    for r, name, dref, nref in layers:
        wc.cell(r, 1, name)
        blue(wc, r, 2, dref, "0.000")
        if isinstance(nref, float):
            blue(wc, r, 3, nref, "0.00")
        else:
            blue(wc, r, 3, nref, "0.00")
        green(wc, r, 4, f"=Parametres!$B$26*B{r}*C{r}", "0.00")
        green(wc, r, 5, f"=Parametres!$B$26*B{r}", "0.00")

    wc["A10"] = "Stockage total (m3)"
    green(wc, 10, 2, "=SUM(D5:D8)", "0.00")
    wc["A11"] = "Écart vs Vcap (m3)"
    green(wc, 11, 2, "=B10-Parametres!B24", "0.00")
    wc["C11"] = "≈ 0 si Af = Vcap/ds"

    ch = BarChart()
    ch.title = "Stockage par couche (m3)"
    ch.add_data(Reference(wc, min_col=4, min_row=4, max_row=8), titles_from_data=True)
    ch.set_categories(Reference(wc, min_col=1, min_row=5, max_row=8))
    wc.add_chart(ch, "A14")

    for col, w in zip("ABCDE", [22, 14, 12, 14, 18]):
        wc.column_dimensions[col].width = w

    # ---------- Scenarios ----------
    ws2 = wb.create_sheet("Scenarios")
    ws2["A1"] = "Scénarios rapides — même bassin, Rd variable"
    ws2["A1"].font = Font(bold=True, size=12, color="1B5E4B")
    ws2["A2"] = (
        "Utile pour montrer au client: 15 / 25 / 40 mm. "
        "Les colonnes C–E sont des formules basées sur Parametres (même C, Ad, ds, f')."
    )
    ws2.merge_cells("A2:F3")

    for j, h in enumerate(["Rd (mm)", "Vcap (m3)", "Af (m2)", "tp (h)", "Vidange OK?"], 1):
        ws2.cell(5, j, h)
    hdr_row(ws2, 5, 5)

    for i, rd in enumerate([15.0, 25.0, 40.0]):
        r = 6 + i
        yellow(ws2, r, 1, rd, "0.0")
        green(ws2, r, 2, f"=Parametres!$B$9*Parametres!$B$7*(A{r}/1000)*Parametres!$B$6", "0.00")
        green(ws2, r, 3, f"=IF(Parametres!$B$25<=0,0,B{r}/Parametres!$B$25)", "0.0")
        green(ws2, r, 4, f"=IF(Parametres!$B$19<=0,0,Parametres!$B$12/(Parametres!$B$19/1000))", "0.0")
        green(
            ws2,
            r,
            5,
            f'=IF(D{r}<=Parametres!$B$20,"OUI","NON")',
        )

    ws2["A11"] = "Note: tp dépend de dp et f', pas de Rd — colonne D identique par scénario."
    for col, w in zip("ABCDE", [12, 12, 12, 10, 12]):
        ws2.column_dimensions[col].width = w

    # ---------- Methode ----------
    wm = wb.create_sheet("Methode")
    wm["A1"] = "Méthode — Bio-rétention v1"
    wm["A1"].font = Font(bold=True, size=13, color="1B5E4B")
    lines = [
        "",
        "Objectif commercial: pack Excel prêt à éditer (jaune) pour pré-dimensionner une cellule de bio-rétention.",
        "",
        "Formules:",
        "  Vcap = η × C × (Rd/1000) × Ad",
        "  ds   = dp + dm·nm + df·nf + dr·nr",
        "  Af   = Vcap / ds",
        "  tp   = dp / (f'/1000)   avec f' en mm/h",
        "",
        "Hypothèses simplifiées (à déclarer dans la note de calcul):",
        "  • Volume de captage type « qualité d'eau / petite pluie » (Rd souvent 25 mm).",
        "  • Stockage = ennoiement + vides média + vides réservoir (porosités effectives).",
        "  • Vidange de surface basée sur f' de conception (essai in situ ÷ facteur de sécurité).",
        "  • Sous-drain: case B21 informative; un design CSA W200 détaille diamètre, pente, IWS, etc.",
        "",
        "Références à citer (ne pas copier le texte normatif):",
        "  • CSA W200:18 (R2023) Design of bioretention systems",
        "  • STEP / Sustainable Technologies — Bioretention: Sizing (guide LID)",
        "  • Exigences municipales locales (Montréal / arrondissement / schéma)",
        "",
        "Prochaines versions possibles:",
        "  v1.1 hydrogramme + bypass",
        "  v1.2 prétraitement (fosse à sédiments)",
        "  v1.3 plantation / sel de voirie (climat froid QC)",
        "",
        "Regénérer: python3 build_bioretention_excel.py",
    ]
    for i, line in enumerate(lines, start=2):
        wm.cell(i, 1, line)
    wm.column_dimensions["A"].width = 100

    # ---------- Checklist ----------
    wk = wb.create_sheet("Checklist")
    wk["A1"] = "Checklist conception (rappels — vérifier CSA W200 + ville)"
    wk["A1"].font = Font(bold=True, size=12, color="1B5E4B")
    items = [
        "Essai d'infiltration in situ → f' de conception documenté",
        "Prétraitement (sédimentation) si voirie / abrasifs d'hiver",
        "Trop-plein / bypass vers réseau pluvial existant",
        "Sous-drain: diamètre, pente, enrobage, regard de nettoyage",
        "Membrane / liner si biofiltre ou protection de fondations",
        "Media mix (sable/fines/MO) adapté drainage vs traitement",
        "Profondeur média vs type de végétation (herbe / arbuste / arbre)",
        "Ennoiement max compatible sécurité / plantation",
        "Temps de vidange ≤ exigence locale (ex. 24–48 h)",
        "Accès entretien (paillis, décompression, vacuum)",
        "Sel de déglaçage / espèces tolérantes (climat froid)",
        "Raccordement légal au réseau / permis / limite de travaux",
    ]
    wk.cell(3, 1, "Item")
    wk.cell(3, 2, "OK? (jaune)")
    wk.cell(3, 3, "Note")
    hdr_row(wk, 3, 3)
    for i, item in enumerate(items):
        r = 4 + i
        wk.cell(r, 1, item)
        yellow(wk, r, 2, "à faire")
        wk.cell(r, 3, "")
        for c in range(1, 4):
            wk.cell(r, c).border = THIN
    for col, w in zip("ABC", [70, 14, 30]):
        wk.column_dimensions[col].width = w

    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    p = build()
    print(f"Wrote {p}")
